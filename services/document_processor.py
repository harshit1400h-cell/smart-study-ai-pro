import io
import re
from dataclasses import dataclass
from pathlib import Path

from config import MAX_UPLOAD_BYTES, SUPPORTED_EXTENSIONS


class DocumentProcessingError(Exception):
    """A supported document could not be safely read or processed."""


@dataclass
class ProcessedDocument:
    text: str
    pages: int
    ocr_used: bool = False
    ocr_warning: str = ""


def validate_upload(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower().lstrip(".")
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentProcessingError("This file type is not supported.")
    if not content:
        raise DocumentProcessingError("The uploaded file is empty.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise DocumentProcessingError("Files must be 20 MB or smaller.")
    return extension


def _ocr_image(image) -> tuple[str, str]:
    try:
        import pytesseract
        from pytesseract import Output
    except ImportError as exc:
        raise DocumentProcessingError(
            "OCR support is unavailable. Install the project dependencies and Tesseract OCR."
        ) from exc

    try:
        text = pytesseract.image_to_string(image).strip()
        data = pytesseract.image_to_data(image, output_type=Output.DICT)
        confidences = [
            float(value) for value in data.get("conf", [])
            if str(value).strip() not in {"", "-1"}
        ]
    except Exception as exc:
        raise DocumentProcessingError(
            "OCR could not read this image. Check that Tesseract OCR is installed and try a clearer image."
        ) from exc

    warning = ""
    if not text:
        raise DocumentProcessingError("No readable text was found in the image.")
    if confidences and sum(confidences) / len(confidences) < 55:
        warning = "Some handwriting or text may have been recognized inaccurately."
    return text, warning


def process_document(filename: str, content: bytes) -> ProcessedDocument:
    extension = validate_upload(filename, content)
    if extension == "txt":
        try:
            text = content.decode("utf-8-sig").strip()
        except UnicodeDecodeError as exc:
            raise DocumentProcessingError("The text file must use UTF-8 encoding.") from exc
        if not text:
            raise DocumentProcessingError("No readable text was found in this file.")
        return ProcessedDocument(text=text, pages=1)

    if extension == "pdf":
        try:
            import pymupdf
            document = pymupdf.open(stream=content, filetype="pdf")
        except Exception as exc:
            raise DocumentProcessingError("This PDF is damaged or could not be opened.") from exc
        if document.page_count == 0:
            raise DocumentProcessingError("This PDF contains no pages.")

        extracted_pages = []
        ocr_used = False
        warnings = []
        for page_number, page in enumerate(document, start=1):
            page_text = page.get_text("text").strip()
            if not page_text:
                try:
                    pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
                    from PIL import Image
                    image = Image.open(io.BytesIO(pixmap.tobytes("png")))
                    page_text, warning = _ocr_image(image)
                    ocr_used = True
                    if warning:
                        warnings.append(warning)
                except DocumentProcessingError as exc:
                    if "No readable text" not in str(exc):
                        raise
            if page_text:
                extracted_pages.append(f"[Page {page_number}]\n{page_text}")
        page_count = document.page_count
        document.close()
        combined = "\n\n".join(extracted_pages).strip()
        if not combined:
            raise DocumentProcessingError(
                "No text could be extracted. For scanned PDFs, install Tesseract OCR and try again."
            )
        return ProcessedDocument(
            text=combined,
            pages=page_count,
            ocr_used=ocr_used,
            ocr_warning=" ".join(dict.fromkeys(warnings)),
        )

    if extension == "docx":
        try:
            from docx import Document
            document = Document(io.BytesIO(content))
            text = "\n".join(p.text for p in document.paragraphs if p.text.strip()).strip()
        except Exception as exc:
            raise DocumentProcessingError("This DOCX file is damaged or could not be opened.") from exc
        if not text:
            raise DocumentProcessingError("No readable text was found in this document.")
        return ProcessedDocument(text=text, pages=1)

    try:
        from PIL import Image, UnidentifiedImageError
        image = Image.open(io.BytesIO(content)).convert("RGB")
    except (ImportError, UnidentifiedImageError, OSError) as exc:
        raise DocumentProcessingError("This image is damaged or could not be opened.") from exc
    text, warning = _ocr_image(image)
    return ProcessedDocument(text=text, pages=1, ocr_used=True, ocr_warning=warning)


def split_into_chunks(text: str, chunk_size: int = 900, overlap: int = 150) -> list[str]:
    if chunk_size <= 0 or overlap < 0 or overlap >= chunk_size:
        raise ValueError("Chunk size must be positive and overlap must be smaller than chunk size.")
    words = re.findall(r"\S+", text)
    if not words:
        return []
    chunks = []
    start = 0
    while start < len(words):
        end = min(start + chunk_size, len(words))
        chunks.append(" ".join(words[start:end]))
        if end == len(words):
            break
        start = end - overlap
    return chunks
